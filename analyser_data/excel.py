"""Lecture des classeurs Excel (.xlsx, .xlsm, .xls) : liste des feuilles et valeurs (issue #10).

Bibliothèques : ``openpyxl`` (.xlsx) et ``xlrd`` (.xls), toutes deux incluses dans Anaconda.
Si ``openpyxl`` manque, un lecteur intégré (bibliothèque standard Python : un .xlsx est un
zip de fichiers XML) prend le relais ; il reconnaît aussi les dates Excel. Le format .xls
(binaire, avant 2007) demande ``xlrd``.

Les lignes lues sont ensuite écrites en CSV temporaire et passent par le même import que
les fichiers texte : détection des en-têtes, du temps, des unités, des voies Graphtec…
"""
from __future__ import annotations

import datetime as dt
import re
import zipfile
from pathlib import Path
from typing import Callable, Dict, Iterator, List, Optional
from xml.etree import ElementTree

from .i18n import _

EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xls"}
_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_PKG_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
# Formats numériques Excel intégrés qui sont des dates / heures
_BUILTIN_DATE_FORMATS = set(range(14, 23)) | {45, 46, 47}
_DATE_CODE_RE = re.compile(r"[dmyhs]", re.IGNORECASE)


class ExcelError(Exception):
    """Classeur illisible (message destiné à l'utilisateur)."""


def is_excel(path) -> bool:
    return Path(path).suffix.lower() in EXCEL_SUFFIXES


def sheet_names(path) -> List[str]:
    """Noms des feuilles du classeur, dans l'ordre."""
    path = Path(path)
    if path.suffix.lower() == ".xls":
        xlrd = _xlrd()
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            return list(book.sheet_names())
        finally:
            book.release_resources()
    try:
        import openpyxl

        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return list(book.sheetnames)
        finally:
            book.close()
    except ImportError:
        return [name for name, _target in _XlsxReader(path).sheets]
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ExcelError(_("Le classeur « {} » est illisible : {}").format(path.name, exc)) from exc


def read_rows(path, sheet: Optional[str] = None, progress: Optional[Callable[[str], None]] = None) -> Iterator[list]:
    """Valeurs de la feuille, ligne par ligne (nombres, textes, dates / heures Python ou None)."""
    path = Path(path)
    names = sheet_names(path)
    if not names:
        raise ExcelError(_("Le classeur « {} » ne contient aucune feuille.").format(path.name))
    sheet = sheet or names[0]
    if sheet not in names:
        raise ExcelError(_("Feuille « {} » introuvable dans « {} ».").format(sheet, path.name))
    if progress:
        progress(_("Lecture de la feuille « {} » de {}…").format(sheet, path.name))
    if path.suffix.lower() == ".xls":
        return _xls_rows(path, sheet)
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return _XlsxReader(path).rows(sheet)
    return _openpyxl_rows(path, sheet)


def _openpyxl_rows(path: Path, sheet: str) -> Iterator[list]:
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for row in book[sheet].iter_rows(values_only=True):
            yield list(row)
    finally:
        book.close()


def _xlrd():
    try:
        import xlrd
    except ImportError as exc:
        raise ExcelError(_("Les fichiers .xls (Excel avant 2007) demandent le module « xlrd », inclus dans "
                           "Anaconda (sinon : conda install xlrd). Ou enregistrez le classeur en .xlsx.")) from exc
    return xlrd


def _xls_rows(path: Path, sheet: str) -> Iterator[list]:
    xlrd = _xlrd()
    book = xlrd.open_workbook(str(path), on_demand=True)
    try:
        data = book.sheet_by_name(sheet)
        for r in range(data.nrows):
            row = []
            for cell in data.row(r):
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    row.append(None)
                else:
                    row.append(cell.value)
            yield row
    finally:
        book.release_resources()


class _XlsxReader:
    """Lecteur .xlsx sans dépendance (zip + XML), utilisé si openpyxl n'est pas installé."""

    def __init__(self, path: Path):
        self.path = path
        try:
            self.zip = zipfile.ZipFile(path)
            book = ElementTree.fromstring(self.zip.read("xl/workbook.xml"))
            rels = ElementTree.fromstring(self.zip.read("xl/_rels/workbook.xml.rels"))
        except (zipfile.BadZipFile, KeyError, ElementTree.ParseError, OSError) as exc:
            raise ExcelError(_("Le classeur « {} » est illisible : {}").format(path.name, exc)) from exc
        targets = {r.get("Id"): r.get("Target") for r in rels.iter(_PKG_REL + "Relationship")}
        self.sheets = []
        for s in book.iter(_MAIN + "sheet"):
            target = targets.get(s.get(_REL + "id"), "")
            target = target.lstrip("/")
            self.sheets.append((s.get("name"), target if target.startswith("xl/") else "xl/" + target))
        pr = book.find(_MAIN + "workbookPr")
        self.epoch = dt.datetime(1904, 1, 1) if pr is not None and pr.get("date1904") in ("1", "true") \
            else dt.datetime(1899, 12, 30)
        self.strings = self._shared_strings()
        self.date_styles = self._date_styles()

    def _shared_strings(self) -> List[str]:
        try:
            root = ElementTree.fromstring(self.zip.read("xl/sharedStrings.xml"))
        except KeyError:
            return []
        return ["".join(t.text or "" for t in si.iter(_MAIN + "t")) for si in root.iter(_MAIN + "si")]

    def _date_styles(self) -> set:
        try:
            root = ElementTree.fromstring(self.zip.read("xl/styles.xml"))
        except KeyError:
            return set()
        custom: Dict[int, str] = {int(f.get("numFmtId")): f.get("formatCode", "")
                                  for f in root.iter(_MAIN + "numFmt")}
        dates = set()
        xfs = root.find(_MAIN + "cellXfs")
        for i, xf in enumerate(xfs if xfs is not None else []):
            fmt = int(xf.get("numFmtId", 0))
            code = re.sub(r'"[^"]*"|\[[^\]]*\]', "", custom.get(fmt, ""))  # textes et couleurs ignorés
            if fmt in _BUILTIN_DATE_FORMATS or (fmt in custom and _DATE_CODE_RE.search(code)):
                dates.add(i)
        return dates

    def rows(self, sheet: str) -> Iterator[list]:
        target = dict(self.sheets)[sheet]
        with self.zip.open(target) as fh:
            for _event, elem in ElementTree.iterparse(fh):
                if elem.tag != _MAIN + "row":
                    continue
                row: list = []
                for c in elem.iter(_MAIN + "c"):
                    col = _column_index(c.get("r")) if c.get("r") else len(row)
                    row.extend([None] * (col - len(row)))
                    row.append(self._value(c))
                yield row
                elem.clear()

    def _value(self, c):
        kind = c.get("t", "n")
        if kind == "inlineStr":
            return "".join(t.text or "" for t in c.iter(_MAIN + "t"))
        v = c.find(_MAIN + "v")
        if v is None or v.text is None:
            return None
        if kind == "s":
            return self.strings[int(v.text)]
        if kind in ("str", "e"):
            return v.text if kind == "str" else None
        if kind == "b":
            return v.text == "1"
        number = float(v.text)
        if int(c.get("s", 0)) in self.date_styles:
            # nombre de jours en virgule flottante : arrondi à la milliseconde (13.999999 s -> 14 s)
            moment = self.epoch + dt.timedelta(milliseconds=round(number * 86_400_000))
            return moment.time() if number < 1 else moment
        return int(number) if number.is_integer() and abs(number) < 1e15 else number


def _column_index(ref: str) -> int:
    """« B3 » -> 1 (colonne B)."""
    index = 0
    for ch in ref:
        if not ch.isalpha():
            break
        index = index * 26 + (ord(ch.upper()) - 64)
    return index - 1


def format_cell(value) -> str:
    """Valeur Excel -> texte pour le CSV intermédiaire (point décimal, dates ISO)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    # Toujours le même format (millisecondes comprises) : un mélange « 13 » / « 13.999999 »
    # empêcherait de reconnaître la colonne temps
    if isinstance(value, dt.datetime):
        value = value + dt.timedelta(microseconds=500)
        return value.replace(microsecond=value.microsecond // 1000 * 1000).isoformat(sep=" ", timespec="milliseconds")
    if isinstance(value, dt.time):
        return value.isoformat(timespec="milliseconds")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        return repr(value)
    return str(value)
